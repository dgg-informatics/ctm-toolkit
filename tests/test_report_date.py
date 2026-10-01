"""report_date — required on every report, and the tie-breaker between reports.

Every genomic doc carries its report's date as REPORT_DATE. Among genomic docs
sharing SAMPLE_ID + TRUE_HUGO_SYMBOL + VARIANT_CATEGORY + TRUE_PROTEIN_CHANGE,
only those from the most recent report go to matching (same-date docs are all
kept). The older rows remain in patient_data, so nothing is lost.
"""
import logging
from datetime import date, datetime

import openpyxl
import pytest
from pydantic import ValidationError

from ctm.schemas.raw.models import RawReportMetadata
from ctm.schemas.raw.normalized import Finding, Patient
from ctm.transformers.excel_reader import read_and_normalize
from ctm.transformers.normalize_manual import normalize_report_metadata
from ctm.transformers.to_matchminer import latest_genomic_docs, to_genomic_docs

PT = "pt_1000000"


def _finding(report_uuid, report_date, biomarker="MET", category="MUTATION", **kw):
    return Finding(pt_uuid=PT, report_uuid=report_uuid, source="x",
                   report_date=report_date, biomarker=biomarker,
                   variant_category=category, **kw)


# ── Required on the report ─────────────────────────────────────────────────────

def test_report_metadata_without_report_date_is_rejected():
    with pytest.raises(ValidationError, match="report_date"):
        RawReportMetadata.model_validate(
            {"report_uuid": "rp_1", "pt_uuid": PT, "source": "tempus"})


def test_report_date_from_excel_is_stored_as_a_plain_date():
    row = RawReportMetadata.model_validate({
        "report_uuid": "rp_1", "pt_uuid": PT, "source": "tempus",
        "report_date": datetime(2026, 9, 1, 0, 0),
    })
    m = normalize_report_metadata(row)
    assert m.report_date == date(2026, 9, 1)
    assert type(m.report_date) is date
    assert "report_date" not in m.raw


def _workbook(tmp_path, report_rows):
    wb = openpyxl.Workbook()
    wb.active.title = "pt_general"
    wb["pt_general"].append(["pt_uuid"])
    wb["pt_general"].append([PT])
    rm = wb.create_sheet("report_metadata")
    rm.append(["report_uuid", "pt_uuid", "source", "report_date"])
    for r in report_rows:
        rm.append(r)
    tf = wb.create_sheet("tempus_findings")
    tf.append(["pt_uuid", "report_uuid", "biomarker", "variant_category", "wildtype"])
    for r in report_rows:
        tf.append([PT, r[0], "MET", "MUTATION", "FALSE"])
    path = tmp_path / "wb.xlsx"
    wb.save(path)
    return path


def test_reader_skips_and_logs_every_report_missing_a_date(tmp_path, caplog):
    path = _workbook(tmp_path, [
        ["rp_1", PT, "tempus", datetime(2026, 9, 1)],
        ["rp_2", PT, "tempus", None],
        ["rp_3", PT, "tempus", "not a date"],
    ])
    with caplog.at_level(logging.ERROR):
        _, metadata, findings = read_and_normalize(path)
    assert [m.report_uuid for m in metadata] == ["rp_1"]
    assert [f.report_uuid for f in findings] == ["rp_1"]   # skipped with its report
    errors = [r.getMessage() for r in caplog.records if r.levelno == logging.ERROR]
    assert len(errors) == 2
    assert "rp_2" in errors[0] and "rp_3" in errors[1]


def test_findings_carry_their_reports_date(tmp_path):
    path = _workbook(tmp_path, [
        ["rp_1", PT, "tempus", datetime(2026, 9, 1)],
        ["rp_2", PT, "tempus", datetime(2026, 6, 1)],
    ])
    _, _, findings = read_and_normalize(path)
    assert {f.report_uuid: f.report_date for f in findings} == {
        "rp_1": date(2026, 9, 1), "rp_2": date(2026, 6, 1)}


# ── Most recent report wins ────────────────────────────────────────────────────

def _latest(*findings):
    return latest_genomic_docs(to_genomic_docs(Patient(pt_uuid=PT), list(findings)))


def test_genomic_docs_carry_their_reports_date():
    docs = to_genomic_docs(Patient(pt_uuid=PT), [_finding("rp_A", date(2026, 9, 1))])
    assert docs[0]["REPORT_DATE"] == "2026-09-01"


def test_most_recent_report_wins_a_conflict():
    """The spec example: Company A (09/01/2026) says MET wildtype, Company B
    (06/01/2026) says MET detected — only A is matched."""
    docs = _latest(_finding("rp_B", date(2026, 6, 1), wildtype=False),
                   _finding("rp_A", date(2026, 9, 1), wildtype=True))
    assert [(d["WILDTYPE"], d["REPORT_DATE"]) for d in docs] == [(True, "2026-09-01")]


def test_same_date_docs_are_all_kept():
    docs = _latest(_finding("rp_A", date(2026, 9, 1), wildtype=True),
                   _finding("rp_B", date(2026, 9, 1), wildtype=False))
    assert len(docs) == 2


# ── Wired through `ctm-mm patients` ────────────────────────────────────────────

def _run_patients(path, out):
    import argparse

    from ctm.mm_cli import _cmd_raw_to_mm
    _cmd_raw_to_mm(argparse.Namespace(excel=str(path), pt_uuid=None, out=str(out)))
    import json
    return json.loads(out.read_text())


def test_patients_command_matches_only_the_newest_report(tmp_path):
    path = _workbook(tmp_path, [
        ["rp_A", PT, "tempus", datetime(2026, 9, 1)],
        ["rp_B", PT, "tempus", datetime(2026, 6, 1)],
    ])
    out = _run_patients(path, tmp_path / "out.json")
    assert len(out["genomic"]) == 1
    assert out["clinical"][0]["REPORT_DATE"] == "2026-09-01"
    reports = {r["report_uuid"]: r for r in out["extras"]["patients"][PT]["reports"]}
    assert reports["rp_A"]["report_date"] == "2026-09-01"
    assert out["genomic"][0]["REPORT_DATE"] == "2026-09-01"
    # The older report's row is not matched, but is still recorded.
    assert len(reports["rp_B"]["findings"]) == 1


def test_patients_command_skips_a_report_missing_its_date(tmp_path):
    path = _workbook(tmp_path, [
        ["rp_A", PT, "tempus", datetime(2026, 6, 1)],
        ["rp_B", PT, "tempus", None],
    ])
    out = _run_patients(path, tmp_path / "out.json")
    assert [d["REPORT_DATE"] for d in out["genomic"]] == ["2026-06-01"]
    reports = out["extras"]["patients"][PT]["reports"]
    assert [r["report_uuid"] for r in reports] == ["rp_A"]
